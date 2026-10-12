"""联调冒烟测试：在 offscreen 模式下验证主窗口与数据层核心功能。

使用临时数据库文件，不写入项目数据；验证完成后清理。
"""

from __future__ import annotations

import os
import shutil
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

sys.path.insert(0, str(Path(__file__).resolve().parent))

from app import config  # noqa: E402

# 使用独立临时存储目录（数据库/图片库/文档库全部隔离，避免污染项目数据）
SMOKE_STORAGE = Path(tempfile.mkdtemp(prefix="smoke_storage_"))
config.set_storage_root(SMOKE_STORAGE)
# 隔离用户设置文件，避免 smoke 写入真实的 app_config.json
config.CONFIG_PATH = SMOKE_STORAGE / "smoke_app_config.json"

from PySide6.QtWidgets import QApplication  # noqa: E402
from qfluentwidgets import isDarkTheme  # noqa: E402

from app.database import Database  # noqa: E402
from app.markdown_utils import extract_title, make_summary, markdown_to_html  # noqa: E402
from app.main_window import MainWindow  # noqa: E402


def check_database() -> None:
    if config.db_path().exists():
        config.db_path().unlink()
    db = Database(config.db_path())
    db.init_schema()

    note_id = db.create_note("# 冒烟测试\n\n**加粗** 与 `代码`", content="# 冒烟测试\n\n**加粗** 与 `代码`")
    assert note_id > 0

    notes = db.list_notes()
    assert len(notes) == 1 and notes[0].id == note_id
    assert notes[0].summary

    note = db.get_note(note_id)
    assert note is not None and "加粗" in (note.content or "")

    db.toggle_favorite(note_id)
    assert db.get_note(note_id).is_favorite is True
    assert len(db.list_notes(favorite_only=True)) == 1

    db.toggle_pinned(note_id)
    assert db.get_note(note_id).is_pinned is True

    db.update_note(note_id, title="更新后的标题", content="更新后的内容")
    note = db.get_note(note_id)
    assert note.title == "更新后的标题" and note.content == "更新后的内容"

    assert len(db.list_notes(keyword="更新")) == 1
    assert len(db.list_notes(keyword="不存在的关键词")) == 0

    db.move_to_trash(note_id)
    assert len(db.list_notes()) == 0
    assert len(db.list_notes(deleted=True)) == 1
    db.restore_note(note_id)
    assert len(db.list_notes()) == 1

    db.move_to_trash(note_id)
    db.delete_permanently(note_id)
    assert len(db.list_notes(deleted=True)) == 0

    db.close()
    print("[ok] 数据库层冒烟通过")


def check_storage_location() -> None:
    """存储位置切换测试：动态路径 + 数据迁移 + 界面重绑定。"""
    import shutil

    from app.app_settings import settings as app_settings

    root = config.storage_root()
    assert config.db_path().parent == root
    assert config.images_dir() == root / "images"
    assert config.documents_dir() == root / "documents"

    new_root = Path(tempfile.mkdtemp(prefix="smoke_new_storage_"))
    (root / "images").mkdir(parents=True, exist_ok=True)
    (root / "documents").mkdir(parents=True, exist_ok=True)
    (root / "images" / "marker.png").write_bytes(b"fake-png")
    (root / "documents" / "keep.docx").write_bytes(b"fake-docx")

    app = QApplication.instance() or QApplication(sys.argv)
    window = MainWindow()
    window.resize(1100, 700)
    window.show()
    app.processEvents()

    window.db.create_note("迁移前笔记", "内容")
    window.refresh_all_pages()
    app.processEvents()

    # 切换：数据迁移 + 连接重建 + 页面重绑定
    assert window.change_storage_location(new_root) is True
    app.processEvents()
    assert (new_root / "notes.db").is_file()
    assert (new_root / "images" / "marker.png").is_file()
    assert (new_root / "documents" / "keep.docx").is_file()
    assert app_settings.storage_dir == str(new_root)
    assert config.storage_root() == new_root.resolve()

    migrated = [n for n in window.db.list_notes() if n.title == "迁移前笔记"]
    assert len(migrated) == 1, [n.title for n in window.db.list_notes()]
    assert window.notes_page.db is window.db
    assert window.trash_page.db is window.db
    window.notes_page.create_note()
    app.processEvents()
    assert len(window.db.list_notes()) == 2

    # 切回原位置（旧数据仍在）
    assert window.change_storage_location(root) is True
    app.processEvents()
    assert any(n.title == "迁移前笔记" for n in window.db.list_notes())

    for note in window.db.list_notes():
        window.db.delete_permanently(note.id)
    window.db.close()
    shutil.rmtree(new_root, ignore_errors=True)
    print("[ok] 存储位置切换通过")


def check_markdown() -> None:
    assert extract_title("# 标题一\n\n正文") == "标题一"
    assert extract_title("第一行就是标题\n第二行") == "第一行就是标题"
    assert extract_title("") == "未命名笔记"
    assert "…" not in make_summary("短内容")
    assert make_summary("# 标题 **粗体**") == "标题 粗体"
    html = markdown_to_html("| a | b |\n|---|---|\n| 1 | 2 |")
    assert "<table>" in html
    html = markdown_to_html("```python\nprint(1)\n```")
    assert "<pre>" in html and "print" in html
    print("[ok] Markdown 工具冒烟通过")


def check_ui() -> None:
    app = QApplication.instance() or QApplication(sys.argv)
    window = MainWindow()
    window.show()
    app.processEvents()

    # 页面与导航
    assert window.notes_page is not None
    assert window.favorites_page is not None
    assert window.trash_page is not None
    assert window.settings_page is not None

    # 新建 -> 选中 -> 编辑 -> 保存
    window.notes_page.create_note()
    app.processEvents()
    note_id = window.notes_page.editor.current_note_id
    assert note_id is not None

    window.notes_page.editor.title_edit.setText("UI 测试标题")
    window.notes_page.editor.edit_area.setPlainText("# UI 测试\n\n正文内容")
    window.notes_page.flush_editor()
    app.processEvents()

    note = window.db.get_note(note_id)
    assert note.title == "UI 测试标题", note.title
    assert note.content == "# UI 测试\n\n正文内容", note.content

    # 视图模式切换
    for mode in ("edit", "preview", "split"):
        window.notes_page.editor._apply_mode(mode)
        app.processEvents()

    # 主题切换
    window.apply_theme("dark")
    app.processEvents()
    assert isDarkTheme() is True
    window.apply_theme("light")
    app.processEvents()

    # 收藏切换与列表刷新
    window.notes_page.editor.favorite_btn.click()
    app.processEvents()
    assert window.db.get_note(note_id).is_favorite is True

    # 回收站恢复流程
    window.notes_page._move_to_trash(note_id)
    app.processEvents()
    assert len(window.db.list_notes(deleted=True)) == 1
    window.trash_page.refresh()
    window.trash_page.note_list.select_note(note_id)
    window.trash_page.restore_selected()
    app.processEvents()
    assert window.db.get_note(note_id).is_deleted is False

    # 搜索
    window.notes_page.note_list.search_edit.setText("UI 测试")
    app.processEvents()
    assert window.notes_page.note_list.selected_note_id() is not None
    window.notes_page.note_list.search_edit.setText("")

    # 全部页面刷新
    window.refresh_all_pages()
    app.processEvents()

    # 清理测试数据
    window.db.delete_permanently(note_id)
    window.db.close()
    print("[ok] UI 冒烟通过")


def check_images() -> None:
    """图片插入功能测试：入库、引用、粘贴、拖放、删除清理。"""
    import shutil

    from PySide6.QtCore import QMimeData, QPointF, Qt, QUrl
    from PySide6.QtGui import QDropEvent, QImage
    from qfluentwidgets import TransparentToolButton  # noqa: F401

    from app import image_utils
    from app.widgets.markdown_editor import MarkdownEditor

    app = QApplication.instance() or QApplication(sys.argv)

    # 1. 外部图片文件复制入库
    work_dir = Path(tempfile.mkdtemp(prefix="smoke_img_work_"))
    src = work_dir / "photo.png"
    image = QImage(48, 32, QImage.Format_RGB32)
    image.fill(Qt.red)
    assert image.save(str(src), "PNG")
    lib_path = image_utils.save_image_file(src)
    assert lib_path.is_file() and lib_path.parent == config.images_dir()

    # 2. Markdown 引用生成与路径提取（含目录安全检查）
    markdown = image_utils.image_markdown(lib_path, alt="测试图")
    assert markdown.startswith("![测试图](images/") and markdown.endswith(".png)")
    paths = image_utils.extract_image_paths(markdown)
    assert len(paths) == 1 and paths[0] == lib_path.resolve() and paths[0].is_file()
    # import_image_markdown 会复制一张新图片并返回其引用
    imported_md = image_utils.import_image_markdown(src, "导入")
    imported_paths = image_utils.extract_image_paths(imported_md)
    assert len(imported_paths) == 1 and imported_paths[0].is_file()
    assert image_utils.remove_images_in_content(imported_md) == 1
    assert not imported_paths[0].exists()
    # 外链与库外路径不应被提取
    assert image_utils.extract_image_paths("![x](https://a.com/b.png)") == []
    assert image_utils.extract_image_paths("![x](../secret.png)") == []

    # 3. 剪贴板 QImage 保存
    clip_lib = image_utils.save_qimage(image)
    assert clip_lib.is_file() and clip_lib.suffix == ".png"

    # 4. 编辑区粘贴图片（模拟 insertFromMimeData）
    editor = MarkdownEditor()
    mime = QMimeData()
    mime.setImageData(image)
    editor.edit_area.insertFromMimeData(mime)
    text = editor.edit_area.toPlainText()
    assert text.count("](images/") == 1, text

    # 5. 拖放图片文件（模拟 dropEvent）
    drop_mime = QMimeData()
    drop_mime.setUrls([QUrl.fromLocalFile(str(src))])
    event = QDropEvent(
        QPointF(10, 10), Qt.CopyAction, drop_mime, Qt.LeftButton, Qt.NoModifier
    )
    editor.edit_area.dropEvent(event)
    text2 = editor.edit_area.toPlainText()
    assert text2.count("](images/") == 2, text2

    # 6. 按内容清理图片（文本引用仍在，但文件应被删除）
    assert image_utils.remove_images_in_content(text2) == 2
    remaining = image_utils.extract_image_paths(text2)
    assert len(remaining) == 2
    assert all(not p.exists() for p in remaining)
    # 6b. markdown 字符串引用的图片同样可被清理
    assert image_utils.remove_images_in_content(markdown) == 1
    assert not lib_path.exists()
    clip_lib.unlink(missing_ok=True)

    # 7. 端到端：彻底删除笔记时数据库层自动清理图片
    db = Database(config.db_path())
    db.init_schema()
    lib2 = image_utils.save_image_file(src)
    note_md = image_utils.image_markdown(lib2, "端到端")
    note_id = db.create_note("图片笔记", note_md)
    db.move_to_trash(note_id)
    db.delete_permanently(note_id)
    assert not lib2.exists(), "彻底删除笔记后图片应被清理"
    # 软删除不清理图片
    lib3 = image_utils.save_image_file(src)
    note_id2 = db.create_note("图片笔记2", image_utils.image_markdown(lib3))
    db.move_to_trash(note_id2)
    assert lib3.exists(), "软删除不应清理图片"
    db.restore_note(note_id2)
    db.delete_permanently(note_id2)
    assert not lib3.exists()
    db.close()

    # 清理临时工作目录与图片库（仅删图片库目录本身）
    shutil.rmtree(work_dir, ignore_errors=True)
    if config.images_dir().exists():
        shutil.rmtree(config.images_dir(), ignore_errors=True)
    print("[ok] 图片插入功能冒烟通过")


def check_documents() -> None:
    """文档导入导出测试：md/txt/docx/pdf -> 笔记 -> docx 导出 + 旧库迁移。"""
    import shutil
    import sqlite3

    from docx import Document as DocxDocument
    from PySide6.QtWidgets import QApplication, QFileDialog, QLabel

    from app import document_utils
    from app.database import Database
    from app.pages.notes_page import NotesPage

    work = Path(tempfile.mkdtemp(prefix="smoke_docs_"))
    config.set_storage_root(work)

    # 1. 准备各格式测试文档
    md_file = work / "note.md"
    md_file.write_text("# 导入的标题\n\n正文内容", encoding="utf-8")
    txt_file = work / "plain.txt"
    txt_file.write_text("纯文本内容", encoding="utf-8")

    docx_file = work / "word.docx"
    doc = DocxDocument()
    doc.add_heading("Word 标题", level=1)
    doc.add_paragraph("普通段落")
    paragraph = doc.add_paragraph("包含")
    bold_run = paragraph.add_run("加粗")
    bold_run.bold = True
    doc.add_paragraph("项目一", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    doc.save(str(docx_file))

    # 异常文档：段落引用不存在的样式（paragraph.style 为 None，曾导致导入崩溃）
    from docx.oxml.ns import qn

    broken_style_file = work / "broken_style.docx"
    broken_doc = DocxDocument()
    broken_doc.add_heading("正常标题", level=1)
    broken_paragraph = broken_doc.add_paragraph("样式缺失的段落")
    p_pr = broken_paragraph._p.get_or_add_pPr()
    p_style = p_pr.makeelement(
        qn("w:pStyle"), {qn("w:val"): "NonExistentStyleId"}
    )
    p_pr.append(p_style)
    broken_doc.add_paragraph("另一个正常段落")
    broken_doc.save(str(broken_style_file))

    pdf_file = work / "doc.pdf"
    import pymupdf

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "PDF page one content", fontsize=12)
    page2 = pdf.new_page()
    page2.insert_text((72, 72), "PDF page two content", fontsize=12)
    pdf.save(str(pdf_file))
    pdf.close()

    # 2. docx -> HTML 富文本转换（保留标题/粗体/列表/表格结构）
    docx_html = document_utils.docx_to_html(docx_file)
    assert "<h1>Word 标题</h1>" in docx_html, docx_html
    assert "font-weight:700" in docx_html  # 粗体
    assert "<li>项目一</li>" in docx_html
    assert "<table" in docx_html and "<td>A</td>" in docx_html
    print("[ok] docx -> HTML 富文本转换通过")

    # 2b. 异常文档（段落样式缺失）不应导致导入失败
    broken_html = document_utils.docx_to_html(broken_style_file)
    assert "样式缺失的段落" in broken_html, broken_html
    assert "正常标题" in broken_html
    assert "另一个正常段落" in broken_html
    print("[ok] 异常样式 docx 防御转换通过")

    # 3. HTML -> docx 导出往返（读回验证结构与格式）
    roundtrip_html = (
        "<h1>导出标题</h1>"
        "<p style=\"text-align:center\">居中的<strong>粗体</strong>与"
        "<em>斜体</em>和<u>下划线</u>，字号"
        "<span style=\"font-size:18pt\">18</span>，颜色"
        "<span style=\"color:#ff0000\">红</span></p>"
        "<ul><li>项目一</li><li>项目二</li></ul>"
        "<table border=\"1\"><tr><td>A</td><td>B</td></tr>"
        "<tr><td>1</td><td>2</td></tr></table>"
    )
    docx_bytes = document_utils.html_to_docx(roundtrip_html, title="导出标题")
    out_docx = work / "out.docx"
    out_docx.write_bytes(docx_bytes)
    check_doc = DocxDocument(str(out_docx))
    all_text = "\n".join(p.text for p in check_doc.paragraphs)
    assert "导出标题" in all_text
    assert "粗体" in all_text and "斜体" in all_text and "下划线" in all_text
    assert "项目一" in all_text
    assert len(check_doc.tables) == 1
    assert check_doc.tables[0].cell(1, 1).text == "2"
    # 格式校验：存在加粗 run
    has_bold = any(
        run.bold
        for para in check_doc.paragraphs
        for run in para.runs
        if run.bold
    )
    assert has_bold, "导出的 docx 应包含加粗文本"
    print("[ok] HTML -> docx 导出往返通过")

    # 4. 旧库迁移：无 source/content_type 列的数据库自动补列
    legacy_db = work / "legacy.db"
    conn = sqlite3.connect(str(legacy_db))
    conn.execute(
        """CREATE TABLE notes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL DEFAULT '未命名笔记',
            content TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            is_pinned INTEGER NOT NULL DEFAULT 0,
            is_favorite INTEGER NOT NULL DEFAULT 0,
            is_deleted INTEGER NOT NULL DEFAULT 0
        )"""
    )
    conn.execute(
        "INSERT INTO notes (title, content, created_at, updated_at)"
        " VALUES ('旧笔记', 'x', '2020-01-01', '2020-01-01')"
    )
    conn.commit()
    conn.close()

    legacy = Database(legacy_db)
    legacy.init_schema()
    old_note = legacy.get_note(1)
    assert old_note is not None and old_note.title == "旧笔记"
    assert old_note.content_type == "md"
    new_id = legacy.create_note(
        "新笔记",
        "内容",
        source_path="documents/x.md",
        source_type="md",
        content_type="md",
    )
    new_note = legacy.get_note(new_id)
    assert new_note.source_path == "documents/x.md"
    assert new_note.source_type == "md"
    assert new_note.content_type == "md"
    legacy.close()
    print("[ok] 数据库 schema 迁移（source/content_type）通过")

    # 5. UI 导入流程（模拟文件对话框多选：md + docx + pdf 三种类型）
    app = QApplication.instance() or QApplication(sys.argv)
    ui_db = Database(config.db_path())
    ui_db.init_schema()
    notes_page = NotesPage(db=ui_db)
    notes_page.resize(1200, 700)
    notes_page.show()
    app.processEvents()

    original_get_names = QFileDialog.getOpenFileNames
    QFileDialog.getOpenFileNames = staticmethod(
        lambda *args, **kwargs: (
            [str(md_file), str(docx_file), str(pdf_file)],
            "",
        )
    )
    try:
        notes_page.import_documents()
    finally:
        QFileDialog.getOpenFileNames = original_get_names
    app.processEvents()

    imported = [
        ui_db.get_note(note.id)
        for note in ui_db.list_notes()
        if note.source_type is not None
    ]
    assert len(imported) == 3, [n.title for n in imported]
    by_type = {note.content_type: note for note in imported}
    assert set(by_type) == {"md", "docx", "pdf"}, set(by_type)
    assert "正文内容" in (by_type["md"].content or "")
    assert "<h1>Word 标题</h1>" in (by_type["docx"].content or "")
    print("[ok] UI 导入流程（md/docx/pdf 三类型）通过")

    # 6. 编辑器路由：按内容类型切换对应编辑器
    notes_page.note_list.select_note(by_type["md"].id)
    app.processEvents()
    assert notes_page.editor is notes_page.markdown_editor
    assert notes_page.editor.current_note_id == by_type["md"].id

    notes_page.note_list.select_note(by_type["docx"].id)
    app.processEvents()
    assert notes_page.editor is notes_page.rich_editor
    rich_html = notes_page.rich_editor.text_edit.toHtml()
    assert "Word 标题" in rich_html, "docx 笔记应在富文本编辑器中渲染"

    notes_page.note_list.select_note(by_type["pdf"].id)
    app.processEvents()
    assert notes_page.editor is notes_page.pdf_viewer
    # PDF 查看器应渲染出页面（第一页 label 有 pixmap）
    page_labels = notes_page.pdf_viewer.container.findChildren(QLabel, "pdfPage")
    assert len(page_labels) == 2, len(page_labels)
    assert len(notes_page.pdf_viewer._page_labels) == 2
    assert page_labels[0].pixmap() is not None, "PDF 第一页应已渲染"
    assert notes_page.pdf_viewer.page_label.text() == "1 / 2"
    print("[ok] 编辑器路由与 PDF 渲染通过")

    # 7. 富文本编辑器编辑往返：重新选中文档 -> 修改 -> flush -> 读回
    notes_page.note_list.select_note(by_type["docx"].id)
    app.processEvents()
    assert notes_page.editor is notes_page.rich_editor
    notes_page.rich_editor.title_edit.setText("改后的标题")
    notes_page.rich_editor.text_edit.setPlainText("编辑后的内容")
    notes_page.flush_editor()
    app.processEvents()
    saved = ui_db.get_note(by_type["docx"].id)
    assert saved.title == "改后的标题"
    assert "编辑后的内容" in (saved.content or "")
    print("[ok] 富文本编辑保存往返通过")

    # 清理测试数据
    for note in imported:
        ui_db.delete_permanently(note.id)
    ui_db.close()
    shutil.rmtree(work, ignore_errors=True)
    if config.images_dir().exists():
        shutil.rmtree(config.images_dir(), ignore_errors=True)
    # 恢复默认存储目录（后续 check_ui 使用）
    config.set_storage_root(SMOKE_STORAGE)
    print("[ok] 文档功能冒烟通过")


def main() -> None:
    check_database()
    check_markdown()
    check_images()
    check_documents()
    check_storage_location()
    check_ui()
    if config.db_path().exists():
        config.db_path().unlink()
    shutil.rmtree(SMOKE_STORAGE, ignore_errors=True)
    print("ALL SMOKE TESTS PASSED")


if __name__ == "__main__":
    main()
